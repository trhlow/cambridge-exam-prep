"""Writes BienBan_ThiThu_Mau.docx, the blank minutes of the mock exam (VE-42): who took part, the 16 situations of
KICH_BAN_THI_THU.md with a result column, the errors found, and the signatures. Times New Roman 13 pt, as the
faculty asks for documents.

    python setup/mock-exam/make_minutes.py   (python-docx)
"""

from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).parent

SITUATIONS = [
    "Thí sinh đăng nhập bằng phiếu và mã PIN",
    "Nhập sai mã PIN",
    "Bắt đầu trước khi giám thị Announce start (cảnh báo StartedEarly)",
    "Làm bài, ghi chú/highlight, đổi cỡ chữ và độ tương phản",
    "Tải lại trang giữa giờ",
    "Tạm dừng một thí sinh, cả phòng; cho làm tiếp",
    "Thông báo cho cả phòng và cho một thí sinh",
    "Mất mạng 1 phút rồi vào lại; cộng giờ bù",
    "Chuyển sang máy khác",
    "Cộng giờ bù sự cố; Edit end time (End in)",
    "Listening: tua lại 1 phút và cộng 1 phút",
    "B1 Writing Part 2: chọn một trong hai câu",
    "End now cho một thí sinh",
    "Hết giờ tự nộp (cả thí sinh thêm 25 % giờ, máy offline)",
    "Giám khảo chấm Writing, xem kết quả, xuất CSV",
    "Yêu cầu xoá thí sinh và admin khác duyệt",
]


def font(document):
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(0)


def centred(document, text, bold=False, size=13):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return paragraph


def table(document, header, rows, widths):
    grid = document.add_table(rows=1, cols=len(header))
    grid.style = "Table Grid"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(grid.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for values in rows:
        cells = grid.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    for row in grid.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
    return grid


def main():
    document = docx.Document()
    font(document)
    for section in document.sections:
        section.left_margin, section.right_margin = Cm(3), Cm(2)
        section.top_margin, section.bottom_margin = Cm(2), Cm(2)

    centred(document, "TRƯỜNG KỸ THUẬT VÀ CÔNG NGHỆ", bold=True)
    centred(document, "KHOA CÔNG NGHỆ THÔNG TIN", bold=True)
    centred(document, "")
    centred(document, "BIÊN BẢN BUỔI THI THỬ", bold=True, size=15)
    centred(document, "Hệ thống thi tiếng Anh trên máy cho Trung tâm Anh ngữ Victory")
    document.add_paragraph("Thời gian: ……… giờ ……… ngày ……/……/2026        Địa điểm: ……………………………………")
    document.add_paragraph("Kỳ thi: ………………………………………… (A2 Key for Schools / B1 Preliminary for Schools)")
    document.add_paragraph("Máy chủ: ……………………  Số máy thí sinh: ……  Trình duyệt: ……………………")

    document.add_paragraph().add_run("1. Người tham gia").bold = True
    table(document, ["Họ tên", "Vai trò", "Máy / tài khoản"], [["", role, ""] for role in ["Admin", "Admin thứ hai", "Giám thị", "Giám khảo"] + ["Thí sinh"] * 6], [6, 4, 5.5])

    document.add_paragraph().add_run("2. Các tình huống đã thử").bold = True
    table(document, ["#", "Tình huống", "Kết quả (Đạt/Không đạt)", "Ghi chú"], [[str(i), text, "", ""] for i, text in enumerate(SITUATIONS, start=1)], [1, 8, 3, 3.5])

    document.add_paragraph().add_run("3. Lỗi phát hiện").bold = True
    table(document, ["#", "Mô tả và bước tái hiện", "Mức độ (Cao/TB/Thấp)", "Mã Jira"], [[str(i), "", "", ""] for i in range(1, 9)], [1, 9, 3, 2.5])

    document.add_paragraph().add_run("4. Nhận xét chung").bold = True
    for _ in range(4):
        document.add_paragraph("……………………………………………………………………………………………………………")

    signatures = document.add_table(rows=2, cols=2)
    for cell, text in zip(signatures.rows[0].cells, ["Người lập biên bản", "Đại diện người tham gia"]):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run(text).bold = True
    for cell in signatures.rows[1].cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run("(Ký, ghi rõ họ tên)").italic = True

    out = HERE / "BienBan_ThiThu_Mau.docx"
    document.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
